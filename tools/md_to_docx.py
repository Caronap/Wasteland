#!/usr/bin/env python3
"""Convert a Markdown file to a DOCX document.

This is a small project-local converter for engineering notes. It supports the
Markdown features used in docs/*.md: headings, paragraphs, bullet/numbered
lists, blockquotes, fenced code blocks, horizontal rules, inline code/bold/
italic/link text, and simple pipe tables.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

try:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError as exc:  # pragma: no cover - friendly CLI failure
    raise SystemExit(
        "Missing dependency: python-docx\n"
        "Install it with: python -m pip install python-docx"
    ) from exc


INLINE_RE = re.compile(
    r"(\[[^\]]+\]\([^)]+\)|`[^`]+`|\*\*[^*]+\*\*|__[^_]+__|\*[^*\n]+\*|_[^_\n]+_)"
)
HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
UL_RE = re.compile(r"^\s*[-*+]\s+(.+)$")
OL_RE = re.compile(r"^\s*\d+[.)]\s+(.+)$")
TABLE_SEPARATOR_RE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$")


def set_font(run_or_style, font_name: str, east_asia_font: str | None = None) -> None:
    font = run_or_style.font
    font.name = font_name
    element = run_or_style._element
    rpr = element.rPr
    if rpr is None:
        rpr = element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = rpr._add_rFonts()
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)
    if east_asia_font:
        rfonts.set(qn("w:eastAsia"), east_asia_font)


def configure_styles(document: Document) -> None:
    styles = document.styles
    set_font(styles["Normal"], "Calibri", "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)

    for level in range(1, 7):
        style_name = f"Heading {level}"
        if style_name in styles:
            set_font(styles[style_name], "Calibri", "Microsoft YaHei")

    if "CodeBlock" not in styles:
        style = styles.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH)
        set_font(style, "Consolas", "Microsoft YaHei")
        style.font.size = Pt(9)
        style.font.color.rgb = RGBColor(40, 52, 62)
        style.paragraph_format.left_indent = Inches(0.25)
        style.paragraph_format.space_before = Pt(3)
        style.paragraph_format.space_after = Pt(3)

    if "QuoteBlock" not in styles:
        style = styles.add_style("QuoteBlock", WD_STYLE_TYPE.PARAGRAPH)
        set_font(style, "Calibri", "Microsoft YaHei")
        style.font.italic = True
        style.font.color.rgb = RGBColor(82, 94, 105)
        style.paragraph_format.left_indent = Inches(0.25)


def add_inline_runs(paragraph, text: str) -> None:
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        token = match.group(0)
        run_text = token
        bold = False
        italic = False
        code = False

        if token.startswith("[") and "](" in token and token.endswith(")"):
            label, url = token[1:-1].split("](", 1)
            run_text = f"{label} ({url})"
            run = paragraph.add_run(run_text)
            run.font.color.rgb = RGBColor(34, 92, 160)
            run.underline = True
            pos = match.end()
            continue
        if token.startswith("`") and token.endswith("`"):
            run_text = token[1:-1]
            code = True
        elif token.startswith("**") and token.endswith("**"):
            run_text = token[2:-2]
            bold = True
        elif token.startswith("__") and token.endswith("__"):
            run_text = token[2:-2]
            bold = True
        elif token.startswith("*") and token.endswith("*"):
            run_text = token[1:-1]
            italic = True
        elif token.startswith("_") and token.endswith("_"):
            run_text = token[1:-1]
            italic = True

        run = paragraph.add_run(run_text)
        run.bold = bold
        run.italic = italic
        if code:
            set_font(run, "Consolas", "Microsoft YaHei")
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(148, 76, 43)
        pos = match.end()

    if pos < len(text):
        paragraph.add_run(text[pos:])


def is_table_separator(line: str) -> bool:
    return bool(TABLE_SEPARATOR_RE.match(line))


def split_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [cell.strip() for cell in line.split("|")]


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=col_count)
    table.style = "Table Grid"
    for row_data in rows:
        row = table.add_row()
        for index in range(col_count):
            cell_text = row_data[index] if index < len(row_data) else ""
            paragraph = row.cells[index].paragraphs[0]
            add_inline_runs(paragraph, cell_text)


def add_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    add_inline_runs(paragraph, text)


def convert_markdown_to_docx(markdown: str, output_path: Path, title: str | None = None) -> None:
    document = Document()
    configure_styles(document)

    if title:
        document.core_properties.title = title

    lines = markdown.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    index = 0
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_lines
        if paragraph_lines:
            add_paragraph(document, " ".join(line.strip() for line in paragraph_lines))
            paragraph_lines = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped == "":
            flush_paragraph()
            index += 1
            continue

        if stripped.startswith("```") or stripped.startswith("~~~"):
            flush_paragraph()
            fence = stripped[:3]
            index += 1
            code_lines: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith(fence):
                code_lines.append(lines[index])
                index += 1
            for code_line in code_lines or [""]:
                document.add_paragraph(code_line, style="CodeBlock")
            index += 1
            continue

        if index + 1 < len(lines) and "|" in line and is_table_separator(lines[index + 1]):
            flush_paragraph()
            table_rows = [split_table_row(line)]
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                table_rows.append(split_table_row(lines[index]))
                index += 1
            add_table(document, table_rows)
            continue

        heading = HEADING_RE.match(line)
        if heading:
            flush_paragraph()
            level = min(len(heading.group(1)), 6)
            document.add_heading(heading.group(2).strip(), level=level)
            index += 1
            continue

        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", stripped):
            flush_paragraph()
            document.add_paragraph("_" * 40)
            index += 1
            continue

        quote_text = stripped[1:].strip() if stripped.startswith(">") else None
        if quote_text is not None:
            flush_paragraph()
            paragraph = document.add_paragraph(style="QuoteBlock")
            add_inline_runs(paragraph, quote_text)
            index += 1
            continue

        unordered = UL_RE.match(line)
        if unordered:
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_runs(paragraph, unordered.group(1))
            index += 1
            continue

        ordered = OL_RE.match(line)
        if ordered:
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Number")
            add_inline_runs(paragraph, ordered.group(1))
            index += 1
            continue

        paragraph_lines.append(line)
        index += 1

    flush_paragraph()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Convert Markdown to DOCX.")
    parser.add_argument("input", type=Path, help="Input .md file")
    parser.add_argument("output", nargs="?", type=Path, help="Output .docx file")
    parser.add_argument("--title", help="Optional document title metadata")
    args = parser.parse_args()

    input_path = args.input
    output_path = args.output or input_path.with_suffix(".docx")
    markdown = input_path.read_text(encoding="utf-8-sig")
    title = args.title or input_path.stem.replace("_", " ").replace("-", " ")
    convert_markdown_to_docx(markdown, output_path, title)
    print(f"Wrote {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
